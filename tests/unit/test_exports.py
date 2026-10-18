from django.http import HttpResponse

from spinalcure.utils import SpinalCureExport


def test_text_csv_excel_neutralise_formula_injection(tmp_path):
    rows = [['=HYPERLINK("http://evil","x")', '+cmd', '-1+1', '@SUM(A1)', 'Asha', 5]]
    headers = ['a', 'b', 'c', 'd', 'e', 'f']
    text = SpinalCureExport.to_text(rows, headers, lambda r: r)
    assert text.splitlines()[1].split('\t')[:4] == ["'=HYPERLINK(\"http://evil\",\"x\")", "'+cmd", "'-1+1", "'@SUM(A1)"]
    csv_body = SpinalCureExport.to_csv(rows, headers, lambda r: r).content.decode()
    assert "'=HYPERLINK" in csv_body and ',+cmd' not in csv_body
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(SpinalCureExport.to_excel(rows, headers, lambda r: r).content))
    ws = wb.active
    assert ws['A2'].value.startswith("'=") and ws['A2'].data_type == 's'    # a string, not a formula
    assert ws['F2'].value == 5 and ws['E2'].value == 'Asha'                    # normal values untouched


def test_to_text_is_tab_separated_and_single_line_per_row():
    out = SpinalCureExport.to_text([['a\tb', 'line1\nline2']], ['h1', 'h2'], lambda r: r)
    assert out.split('\n') == ['h1\th2', 'a b\tline1 line2']
