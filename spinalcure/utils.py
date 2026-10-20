"""
Utilities for SpinalCure application.
"""
import logging
import csv
import os
import uuid
import base64
import binascii
from datetime import datetime
from io import BytesIO

from rest_framework.views import exception_handler
from rest_framework.response import Response
from rest_framework import status, serializers
from django.conf import settings
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.drawing.image import Image as XLImage

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer, Image as PDFImage
)
from reportlab.graphics.shapes import Drawing, Circle
from reportlab.graphics import renderPDF
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Exception Handler
# ---------------------------------------------------------------------------

def custom_exception_handler(exc, context):
    """
    Custom exception handler for REST API.
    Returns professional error responses without exposing sensitive info.
    """
    response = exception_handler(exc, context)

    if response is None:
        logger.error(f"Unhandled exception in {context['view']}", exc_info=exc)
        err_msg = str(exc) if settings.DEBUG else 'An unexpected error occurred. Please try again later.'
        return Response(
            {'detail': err_msg, 'status': 'error'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )

    # Hardened check for linter and safety
    status_code = getattr(response, 'status_code', 500)
    data = getattr(response, 'data', {})
    logger.warning(f"{status_code} error in {context['view']}: {data}")

    if 'detail' in response.data:
        if isinstance(response.data['detail'], str):
            response.data = {
                'detail': response.data['detail'],
                'status': 'error',
                'code': response.status_code
            }

    return response


# ---------------------------------------------------------------------------
# UUID Utilities
# ---------------------------------------------------------------------------

def encode_b64uuid(uuid_obj):
    """Encodes a UUID object to a URL-safe Base64 string (22 chars)."""
    if isinstance(uuid_obj, str):
        try:
            uuid_obj = uuid.UUID(uuid_obj)
        except ValueError:
            return None
    if not isinstance(uuid_obj, uuid.UUID):
        return None
    return base64.urlsafe_b64encode(uuid_obj.bytes).decode('ascii').rstrip('=')


def decode_b64uuid(b64_str):
    """Decodes a URL-safe Base64 string back to a UUID object."""
    try:
        padding = '=' * (4 - (len(b64_str) % 4))
        bytes_obj = base64.urlsafe_b64decode(b64_str + padding)
        return uuid.UUID(bytes=bytes_obj)
    except (ValueError, TypeError, binascii.Error):
        return None


def parse_uuid_param(value):
    """Accept a base64-url UUID or a canonical UUID string; return a UUID or None if it is neither."""
    if isinstance(value, uuid.UUID):
        return value
    value = (value or '').strip() if isinstance(value, str) else ''
    if not value:
        return None
    if len(value) == 22:
        decoded = decode_b64uuid(value)
        if decoded:
            return decoded
    try:
        return uuid.UUID(value)
    except ValueError:
        return None


class B64UUIDField(serializers.Field):
    """
    Serializes Django UUIDs to Base64 Url-Safe Strings and vice versa.
    """
    def to_representation(self, value):
        if value is None:
            return None
        return encode_b64uuid(value)

    def to_internal_value(self, data):
        if not data:
            return None
        uuid_val = decode_b64uuid(data)
        if uuid_val is None:
            raise serializers.ValidationError('Invalid Base64 UUID format.')
        return uuid_val


# ---------------------------------------------------------------------------
# CircularImage — custom ReportLab flowable
# ---------------------------------------------------------------------------

class CircularImage(PDFImage):
    """Custom Flowable to render a circular clipped image in ReportLab."""
    def draw(self):
        self.canv.saveState()
        path = self.canv.beginPath()
        r = min(self.drawWidth, self.drawHeight) / 2.0
        path.circle(self.drawWidth / 2.0, self.drawHeight / 2.0, r)
        self.canv.clipPath(path, stroke=0)
        PDFImage.draw(self)
        self.canv.restoreState()


# ---------------------------------------------------------------------------
# SpinalCureExport — unified export utility
# image_cols uses 0-BASED column indices throughout all methods
# ---------------------------------------------------------------------------

class SpinalCureExport:
    """Unified utility to export data to Excel, CSV, PDF, and Text formats."""

    @staticmethod
    def _get_timestamp():
        return datetime.now().strftime('%Y%m%d_%H%M')

    @staticmethod
    def _cell(value):
        """Neutralise spreadsheet formula injection: text starting with = + - @ (or tab/CR) is prefixed with '."""
        if isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r'):
            return "'" + value
        return value

    @staticmethod
    def to_text(queryset, headers, mapping_func):
        """Tab-separated text for the "copy to clipboard" export."""
        lines = ['\t'.join(str(h) for h in headers)]
        for item in queryset:
            lines.append('\t'.join(str(SpinalCureExport._cell(v)).replace('\t', ' ').replace('\n', ' ')
                                   for v in mapping_func(item)))
        return '\n'.join(lines)

    @staticmethod
    def to_csv(queryset, headers, mapping_func, filename_prefix="report"):
        """Export queryset or list to CSV."""
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = (
            f'attachment; filename="{filename_prefix}_{SpinalCureExport._get_timestamp()}.csv"'
        )
        writer = csv.writer(response)
        writer.writerow(headers)
        for item in queryset:
            writer.writerow([SpinalCureExport._cell(v) for v in mapping_func(item)])
        return response

    @staticmethod
    def to_excel(queryset, headers, mapping_func, sheet_name="Data",
                 filename_prefix="report", image_cols=None):
        """Export queryset or list to Excel with professional styling.
        image_cols: list of 0-based column indices that contain image paths.
        """
        if image_cols is None:
            image_cols = []

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = (
            f'attachment; filename="{filename_prefix}_{SpinalCureExport._get_timestamp()}.xlsx"'
        )

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_num, item in enumerate(queryset, 2):
            row_data = mapping_func(item)
            for col_idx, value in enumerate(row_data):  # 0-indexed
                col_num = col_idx + 1  # 1-indexed for openpyxl
                cell = ws.cell(row=row_num, column=col_num)

                if col_idx in image_cols:
                    if value:
                        try:
                            img_path = str(value)
                            if not os.path.isabs(img_path):
                                img_path = os.path.join(str(settings.MEDIA_ROOT), img_path)
                            if os.path.exists(img_path):
                                img = XLImage(img_path)
                                img.width, img.height = 40, 40
                                ws.add_image(img, cell.coordinate)
                                ws.row_dimensions[row_num].height = 35
                                cell.value = ""  # Clear the path text
                            else:
                                cell.value = "No Photo"
                        except Exception:
                            cell.value = "Error"
                    else:
                        cell.value = ""
                else:
                    cell.value = SpinalCureExport._cell(value)

                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 50)

        wb.save(response)
        return response

    @staticmethod
    def to_pdf(queryset, headers, mapping_func, title="Report",
               filename_prefix="report", orientation='portrait', image_cols=None):
        """Export queryset or list to PDF with professional table layout.
        image_cols: list of 0-based column indices that contain image file paths.
                    Images are rendered as circular thumbnails.
        """
        if image_cols is None:
            image_cols = []

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = (
            f'attachment; filename="{filename_prefix}_{SpinalCureExport._get_timestamp()}.pdf"'
        )

        buffer = BytesIO()
        pagesize = landscape(A4) if orientation == 'landscape' else A4
        doc = SimpleDocTemplate(
            buffer, pagesize=pagesize,
            rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=20
        )
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle', parent=styles['Heading1'],
            fontSize=16, spaceAfter=4, alignment=1,
            textColor=colors.HexColor("#1f2937")
        )
        meta_style = ParagraphStyle(
            'ReportMeta', fontSize=8,
            textColor=colors.grey, alignment=1, spaceAfter=10
        )
        cell_style = ParagraphStyle(
            'CellStyle', parent=styles['Normal'],
            fontSize=7, leading=9, wordWrap='LTR'
        )
        header_cell_style = ParagraphStyle(
            'HeaderCell', parent=styles['Normal'],
            fontSize=8, leading=10, textColor=colors.white,
            fontName='Helvetica-Bold'
        )

        from .pdftext import para
        elements.append(para(f"SpinalCure — {title}", title_style))
        elements.append(Paragraph(
            f"Generated: {datetime.now().strftime('%B %d, %Y %I:%M %p')}",
            meta_style
        ))
        elements.append(Spacer(1, 0.15 * inch))

        # Build header row
        header_row = [para(h, header_cell_style) for h in headers]
        data = [header_row]

        # Build data rows — image_cols is 0-indexed
        for item in queryset:
            raw_row = mapping_func(item)
            row = []
            for col_idx, val in enumerate(raw_row):  # 0-indexed
                if col_idx in image_cols:
                    if val and isinstance(val, str) and os.path.exists(val):
                        try:
                            img = CircularImage(val, width=36, height=36)
                            row.append(img)
                        except Exception:
                            row.append(para("N/A", cell_style))
                    else:
                        row.append(Paragraph("", cell_style))
                else:
                    row.append(para(val, cell_style))
            data.append(row)

        # Smart column widths
        num_cols = len(headers)
        page_width = pagesize[0] - 40
        col_widths = []
        for h in headers:
            if h in ("Chief Complaint", "Address"):
                col_widths.append(1.8 * (page_width / num_cols))
            elif h in ("Email", "Name"):
                col_widths.append(1.4 * (page_width / num_cols))
            elif h in ("S.No", "Photo"):
                col_widths.append(0.55 * (page_width / num_cols))
            elif h in ("Gender", "Progress"):
                col_widths.append(0.7 * (page_width / num_cols))
            else:
                col_widths.append(page_width / num_cols)

        total = sum(col_widths)
        col_widths = [(w / total) * page_width for w in col_widths]

        table = Table(data, colWidths=col_widths, repeatRows=1)

        style = TableStyle([
            # Header
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            # Body rows
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#d1d5db")),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 1), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
            # Alternating row colors
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
        ])

        # Center-align image columns
        for col_idx in image_cols:
            style.add('ALIGN', (col_idx, 0), (col_idx, -1), 'CENTER')
            style.add('VALIGN', (col_idx, 0), (col_idx, -1), 'MIDDLE')

        table.setStyle(style)
        elements.append(table)

        doc.build(elements)
        pdf = buffer.getvalue()
        buffer.close()
        response.write(pdf)
        return response

        output = "\t".join(headers) + "\n"
        for item in queryset:
            row = [str(val) for val in mapping_func(item)]
            output += "\t".join(row) + "\n"
        return output


def log_action(request, action, model_name, object_repr, object_id='', changes=None, raw_changes=None, snapshot=None):
    """
    Create an AuditLog entry for a CRUD action.

    Args:
        request: Django HttpRequest (for user and IP)
        action: 'create', 'update', or 'delete'
        model_name: String name of the model/module, e.g. 'Patient'
        object_repr: Human-readable description of the record, e.g. 'Kavitha Balakrishnan'
        object_id: The primary key of the record (optional)
        changes: dict of {"field": {"old": "old_value", "new": "new_value"}} (optional)
        raw_changes: dict of raw database field changes for programmatic reverting (optional)
        snapshot: complete serialized JSON dict of the object before deletion (optional)
    """
    try:
        from clinic.models import AuditLog
        ip = request.META.get('HTTP_X_FORWARDED_FOR', '').split(',')[0].strip() or \
             request.META.get('REMOTE_ADDR')
        log_entry = AuditLog.objects.create(
            user=request.user if request.user.is_authenticated else None,
            action=action,
            model_name=model_name,
            object_id=str(object_id),
            object_repr=str(object_repr)[:300],
            changes=changes,
            raw_changes=raw_changes,
            snapshot=snapshot,
            ip_address=ip or None,
        )
        
        # Optionally send an email notification to superadmin for sensitive actions
        if request.user and getattr(request.user, 'role', '') != 'superadmin' and not getattr(request.user, 'is_superuser', False):
            if action in ['update', 'delete']:
                try:
                    from accounts.models import CustomUser
                    from django.core.mail import EmailMultiAlternatives
                    from django.template.loader import render_to_string
                    from django.conf import settings
                    from django.utils import timezone
                    
                    superadmin = CustomUser.objects.filter(is_superuser=True).first()
                    if not superadmin:
                        superadmin = CustomUser.objects.filter(role='superadmin').first()
                        
                    admin_email = superadmin.email if superadmin else getattr(settings, 'DEFAULT_FROM_EMAIL', 'admin@spinalcure.com')
                    
                    context = {
                        'staff_name': request.user.get_full_name() or request.user.email,
                        'staff_role': getattr(request.user, 'get_role_display', lambda: 'Staff')(),
                        'staff_email': request.user.email,
                        'action_type': action.upper(),
                        'target_model': model_name,
                        'target_name': object_repr,
                        'timestamp': timezone.now().strftime('%d %b %Y, %I:%M %p'),
                        'changes': changes or {}
                    }
                    
                    subject = f'SpinalCure Audit Log: {action.capitalize()} {model_name}'
                    text_message = f"Audit Log:\n{context['staff_name']} {action}d {model_name}: {object_repr}"
                    html_message = render_to_string('emails/audit_log_notification.html', context, request=request)
                    
                    msg = EmailMultiAlternatives(
                        subject=subject,
                        body=text_message,
                        from_email=getattr(settings, 'DEFAULT_FROM_EMAIL', 'noreply@spinalcure.com'),
                        to=[admin_email],
                    )
                    msg.attach_alternative(html_message, "text/html")
                    msg.send(fail_silently=True)
                except Exception as mail_err:
                    logger.error(f"Failed to send audit log email: {mail_err}")

    except Exception as e:
        logger.warning(f'AuditLog failed to record action [{action}] on [{model_name}]: {e}')


def safe_next(request, fallback):
    """A `next` value from the request, only if it points back into this site; otherwise `fallback`."""
    from django.utils.http import url_has_allowed_host_and_scheme
    target = request.POST.get('next') or request.GET.get('next') or ''
    if target and url_has_allowed_host_and_scheme(target, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return target
    return fallback
